from docx import Document
import os

doc_path = r'g:\منطقة الحرة السويس - 2025\تموينات السفن\بدون مشروعات غير مهمه - دراسة كاملة - لتطوير قطاع تموين السفن في المناطق الحرة لقناة السويس.docx'

doc = Document(doc_path)

print('=== DOCUMENT STRUCTURE ===')
print(f'Total paragraphs: {len(doc.paragraphs)}')
print(f'Total tables: {len(doc.tables)}')
print()

# Extract all text to a file for analysis
output_path = r'g:\منطقة الحرة السويس - 2025\تموينات السفن\document_content.txt'
with open(output_path, 'w', encoding='utf-8') as f:
    f.write('=== FULL DOCUMENT TEXT ===\n\n')
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            f.write(f'[P{i}] {p.text}\n\n')
    
    f.write('\n\n=== TABLES ===\n\n')
    for ti, table in enumerate(doc.tables):
        f.write(f'\n--- Table {ti+1} ---\n')
        for row in table.rows:
            row_text = ' | '.join([cell.text.strip() for cell in row.cells])
            f.write(row_text + '\n')

print(f'Content extracted to: {output_path}')
print(f'File size: {os.path.getsize(output_path)} bytes')
